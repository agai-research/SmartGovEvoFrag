"""Write docs/Dataset.tex and docs/Dataset.docx: dataset summary and the recorded-results tables (python docs/build_dataset.py)."""
import collections
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from docs.build_results import SECTIONS, SIGCOLS, SIGHDR, TABLES, sig_rows
from docs.rlib import read_csv
from docs.tabs import add_docx_table, tex_table

M = ROOT / "Data" / "main"
INST = ROOT / "Data" / "instances"


def jl(p):
    return json.loads((M / p).read_text())


def base_stats():
    man = jl("manifest.json")
    fr = jl("fragments.json")
    kinds = collections.Counter(f["kind"] for f in fr)
    ftypes = collections.Counter(f["ftype"] for f in fr)
    pr = jl("procedures.json")
    pk = collections.Counter(p["kind"] for p in pr)
    seed = {n: len(jl("seed_%s.json" % n)) for n in ("services", "resources", "requirements", "constraints", "fragments", "service_links")}
    return man, kinds, ftypes, pk, seed, pr


def full_files(exp):
    """The instance files of the reported (full-profile) series: five runs, one file per x value."""
    import importlib
    from hfc import config
    cfg = config.load()
    mod = importlib.import_module("Exp." + exp)
    prof = cfg["experiments"]["full"]
    labels = mod.plan(cfg, prof[exp])["labels"]
    if getattr(mod, "SHARED", False):
        labels = labels[:1]
    names = ["run%d_x%s.json" % (r, lab) for r in range(prof["runs"]) for lab in labels]
    return [INST / exp / n for n in names if (INST / exp / n).exists()]


def instance_stats():
    """Per experiment: files, requests, goals, records, events and the repository sizes."""
    out = []
    for d in sorted(INST.iterdir(), key=lambda p: int(p.name[1:]) if p.name[1:].isdigit() else 99):
        if not d.is_dir():
            continue
        files = full_files(d.name)
        nreq = nrec = nev = ng = 0
        sizes = set()
        kinds = collections.Counter()
        for fp in files:
            x = json.loads(fp.read_text())
            nreq += len(x["requests"])
            nrec += len(x["records"])
            nev += len(x["events"])
            ng += sum(len(q["truth"]["goals"]) for q in x["requests"])
            sizes.add(x["size"])
            kinds.update(e["kind"] for e in x["events"])
        out.append((d.name, len(files), nreq, ng, nrec, nev, sorted(sizes), kinds))
    return out


def prompt_stats():
    p = json.loads((ROOT / "Prompts" / "prompts.json").read_text())
    return len(p["queries"])


def tables():
    man, kinds, ftypes, pk, seed, pr = base_stats()
    T = []
    T.append(("tab:ds-base", "Base repository built from the seed dataset.", ["quantity", "value"],
              [["economies", man["economies"]], ["procedures", man["procedures"]], ["hot procedures (requested most often)", man["hot_procedures"]],
               ["atomic fragments", kinds.get("atm", 0)], ["composite fragments in the base", kinds.get("cmp", 0)], ["hyperedges", man["hyperedges"]],
               ["hyperedges SEQ", man["edges_by_pattern"]["seq"]], ["hyperedges AND", man["edges_by_pattern"]["and"]], ["hyperedges XOR", man["edges_by_pattern"]["xor"]],
               ["hyperedges LOOP", man["edges_by_pattern"]["loop"]]]))
    T.append(("tab:ds-seed", "Seed records of the B-READY dataset used for the construction.", ["table", "records"],
              [["services", seed["services"]], ["resources", seed["resources"]], ["requirements", seed["requirements"]], ["constraints", seed["constraints"]],
               ["seed fragments", seed["fragments"]], ["service link groups", seed["service_links"]]]))
    T.append(("tab:ds-kinds", "Procedures by kind and fragments by type.", ["category", "name", "count"],
              [["procedure kind", k, v] for k, v in sorted(pk.items())] + [["fragment type", k, v] for k, v in sorted(ftypes.items())]))
    rows = []
    for name, nf, nreq, ng, nrec, nev, sizes, ek in instance_stats():
        rows.append([name, nf, nreq, ng, nrec, nev, ", ".join(str(s) for s in sizes)])
    T.append(("tab:ds-inst", "Experiment instances: files, requests, goals, stored records, change events and repository sizes.",
              ["series", "instance files", "requests", "goals", "records", "events", "repository sizes |F|"], rows))
    ev = collections.Counter()
    for d in INST.iterdir():
        if d.is_dir():
            for fp in full_files(d.name):
                ev.update(e["kind"] for e in json.loads(fp.read_text())["events"])
    T.append(("tab:ds-events", "Change events over all instances.", ["event kind", "count"], [[k, v] for k, v in sorted(ev.items())]))
    T.append(("tab:ds-prompts", "Test prompts of Prompts/prompts.json.", ["quantity", "value"], [["queries (scenario and variants)", prompt_stats()]]))
    return T


def result_tables():
    """The same tables, in the same order and with the same rows and columns, as Results.tex."""
    out = []
    for _, title, exps in SECTIONS:
        for e in exps:
            for k, (csv, cap) in enumerate(TABLES[e]):
                if not (ROOT / "results" / "tables" / csv).exists():
                    continue
                hdr, rows = sig_rows(csv) if csv.endswith("_significance.csv") else read_csv(csv)
                out.append((e, title, "tab:r-%s-%d" % (e, k + 1), cap, hdr, rows))
    return out


def build():
    T = tables()
    R = result_tables()
    tex = ["% Dataset summary and recorded results; built by docs/build_dataset.py", "\\subsection{Dataset}", "\\label{sec:dataset}", "",
           "The base repository is built from the B-READY business-environment seed dataset: its services, resources, requirements, constraints and seed fragments are expanded to 50 economies and 65 smart-governance procedure families "
           "(3{,}250 procedures) whose steps are atomic fragments connected by SEQ, AND, XOR and LOOP hyperedges. Composite fragments do not exist in the base; they appear when the memory promotes a recurring region. "
           "Requests, stored compositions and change events are built from the repository with fixed seeds, in three disjoint splits (train, validation, test); the reported experiments use the test split.", ""]
    for lab, cap, hdr, rows in T:
        tex += [tex_table(hdr, rows, cap, lab, rules=False), ""]
    tex += ["\\subsubsection{Recorded results}", "The tables below repeat the tables of the results section, with the same rows and columns.", ""]
    for e, title, lab, cap, hdr, rows in R:
        tex += [tex_table(hdr, rows, cap, lab), ""]
    (ROOT / "docs" / "Dataset.tex").write_text("\n".join(tex))

    from docx import Document
    doc = Document()
    doc.add_heading("HyFragComp dataset", 0)
    doc.add_paragraph("Base repository built from the B-READY business-environment seed dataset and the instances built from it for the experiments.")
    doc.add_heading("Dataset summary", 1)
    for lab, cap, hdr, rows in T:
        add_docx_table(doc, hdr, rows, cap)
    doc.add_heading("Recorded results", 1)
    for e, title, lab, cap, hdr, rows in R:
        add_docx_table(doc, hdr, rows, cap, size=6)
    doc.save(ROOT / "docs" / "Dataset.docx")
    return len(T), len(R)


if __name__ == "__main__":
    print("dataset tables, result tables:", build())
